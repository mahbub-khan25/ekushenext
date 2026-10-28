#!/usr/bin/env python3
"""Make the spelling-review workbook: the open spelling questions in review-spelling.tsv, most common
first, with a dropdown for the answer. Fill it in (Excel or Google Sheets), then bring the answers
into curated/spelling-seed.tsv with import_review.py.

Usage: review_sheet.py data/dictionary/review-spelling.tsv OUT.xlsx [ROWS=500]
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ANSWERS = [
    "B is right (A is a misspelling)",
    "A is right (B is a misspelling)",
    "Both right (different words)",
    "Not sure",
]
QUESTIONS = {  # (plain question, Ekushe's guess)
    "less common spelling of the same word?": ("Is A just a rarer spelling of B?", ANSWERS[0]),
    "common misspelling, or a separate word?": ("Is A a misspelling of B, or a different word?", ANSWERS[0]),
    "both common: variant or separate word?": ("Both are common. Same word spelled two ways, or two words?", ""),
    "rule prefers the less common spelling": ("The Bangla Academy rule gives B, though A is more common. Is B right?", ANSWERS[0]),
    "kept ী because usage is overwhelming; Academy may prefer ি": ("Usage strongly favours B. Is B right?", ANSWERS[0]),
}

LATIN = "Arial"
BANGLA = "Nirmala UI"   # Windows' Bangla font; Excel and Sheets fall back to another Bangla font elsewhere
HEAD_FILL = PatternFill("solid", fgColor="1F4E46")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="C9C9C9")


def main():
    src, out = sys.argv[1], sys.argv[2]
    n = int(sys.argv[3]) if len(sys.argv) > 3 else 500
    rows = []
    with open(src, encoding="utf-8") as f:
        r = csv.reader(f, delimiter="\t")
        next(r)
        for x in r:
            if len(x) >= 7 and not x[6].strip():
                rows.append(x)
    rows.sort(key=lambda x: -max(float(x[2] or 0), float(x[3] or 0)))
    rows = rows[:n]

    wb = Workbook()
    how = wb.active
    how.title = "How to fill"
    rv = wb.create_sheet("Review")

    # ---- Review sheet
    heads = ["#", "Word A", "Word B (Ekushe's standard)", "A per million words", "B per million words",
             "Question", "Ekushe's guess", "Your answer", "Note (optional)", "rule"]
    widths = [6, 18, 24, 12, 12, 52, 30, 32, 30, 0]
    for c, (h, w) in enumerate(zip(heads, widths), 1):
        cell = rv.cell(row=1, column=c, value=h)
        cell.font = Font(name=LATIN, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        rv.column_dimensions[cell.column_letter].width = w or 8
    rv.column_dimensions["J"].hidden = True   # the build's rule name, for the importer
    rv.row_dimensions[1].height = 32
    rv.freeze_panes = "B2"

    example = ["Example", "সুরু", "শুরু", 3.01, 1482.36, "Is A a misspelling of B, or a different word?",
               ANSWERS[0], ANSWERS[0], "সুরু is a misspelling", "example"]
    data = [example]
    for i, x in enumerate(rows, 1):
        q, guess = QUESTIONS.get(x[5], (x[5], ""))
        data.append([i, x[0], x[1], round(float(x[2] or 0), 2), round(float(x[3] or 0), 2), q, guess, None, None, x[4]])
    for r_i, row in enumerate(data, 2):
        for c, v in enumerate(row, 1):
            cell = rv.cell(row=r_i, column=c, value=v)
            bangla = c in (2, 3)
            cell.font = Font(name=BANGLA if bangla else LATIN, size=13 if bangla else 10,
                             italic=(r_i == 2), color="555555" if r_i == 2 else "000000")
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="center", wrap_text=c in (6, 9))
            if c in (4, 5):
                cell.number_format = "#,##0.00"
            if r_i == 2:
                cell.fill = EXAMPLE_FILL
            elif c in (8, 9):
                cell.fill = INPUT_FILL
    last = len(data) + 1
    dv = DataValidation(type="list", formula1='"' + ",".join(ANSWERS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Pick an answer",
                        error="Choose one of the four answers from the list.")
    rv.add_data_validation(dv)
    dv.add(f"H3:H{last}")
    rv.auto_filter.ref = f"A1:I{last}"

    # ---- How to fill
    how.column_dimensions["A"].width = 34
    how.column_dimensions["B"].width = 70
    how["A1"] = "Ekushe Next — spelling review"
    how["A1"].font = Font(name=LATIN, bold=True, size=16, color="1F4E46")
    lines = [
        ("What this is", f"The {len(rows)} most common spelling questions Ekushe's dictionary build could not settle by itself, most frequent first. Your answers decide which spelling Ekushe offers first."),
        ("What to fill in", "Only the yellow cells on the Review sheet: 'Your answer' (pick from the list) and, if you like, a note."),
        ("The answers", "B is right (A is a misspelling) · A is right (B is a misspelling) · Both right (different words) · Not sure. Leave a row empty to skip it."),
        ("Ekushe's guess", "What the build would do without you. Agree or not; it is only a hint."),
        ("Standard", "Bangla Academy standard spelling (প্রমিত বানান) is the target, even when another spelling is more common online."),
        ("Row 2", "The grey row is an example of a filled-in answer; it is ignored."),
        ("When done", "Save the file (or download it from Google Sheets as .xlsx or .csv) and send it back."),
    ]
    for i, (k, v) in enumerate(lines, 3):
        how.cell(row=i, column=1, value=k).font = Font(name=LATIN, bold=True)
        c = how.cell(row=i, column=2, value=v)
        c.font = Font(name=LATIN)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    p = len(lines) + 4
    how.cell(row=p, column=1, value="Progress").font = Font(name=LATIN, bold=True, size=12, color="1F4E46")
    stats = [("Questions", f"=COUNT(Review!A3:A{last})")]
    for a in ANSWERS:
        stats.append((a, f'=COUNTIF(Review!H3:H{last},"{a}")'))
    stats.append(("Answered", f"=COUNTA(Review!H3:H{last})"))
    stats.append(("Still open", f"=B{p + 1}-B{p + 1 + len(ANSWERS) + 1}"))
    for i, (k, f) in enumerate(stats, p + 1):
        how.cell(row=i, column=1, value=k).font = Font(name=LATIN)
        c = how.cell(row=i, column=2, value=f)
        c.font = Font(name=LATIN, bold=k in ("Answered", "Still open"))
        c.alignment = Alignment(horizontal="left")
    wb.save(out)
    print(f"{out}: {len(rows)} questions")


if __name__ == "__main__":
    main()
