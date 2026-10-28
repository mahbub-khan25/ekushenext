#!/usr/bin/env python3
"""Turn filled-in typing kits (typing_kit.py) into the held-out test set eval/community.tsv.

Each typed sentence with one typed word per Bangla word becomes one test per word: the typed words up
to it, scored on that word, with the words before it as context. Sentences whose word counts don't
match are reported and skipped. Re-running with the same kits gives the same file; tests from other
kits already in the file are kept.

Usage: import_typing.py eval/community.tsv KIT.xlsx [KIT2.xlsx ...]
"""
import re
import sys
import unicodedata

PUNCT = re.compile(r"[।,?!.;:\"'“”‘’()\-–—…]")


def words(s):
    return [w for w in PUNCT.sub(" ", s or "").split() if w]


def kit_rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    how = wb["How to fill"]
    about = {}
    for row in how.iter_rows(values_only=True):
        if row and row[0] in ("Your first name (optional)", "District you grew up in", "Keyboard you use now"):
            about[row[0]] = (row[1] or "").strip() if isinstance(row[1], str) else ""
    who = " ".join(x for x in (about.get("District you grew up in"), about.get("Keyboard you use now")) if x)
    for row in wb["Sentences"].iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, "Example") or not row[2]:
            continue
        yield str(row[0]), unicodedata.normalize("NFC", row[1]), str(row[2]), who


def main():
    out, kits = sys.argv[1], sys.argv[2:]
    tag = lambda p: re.sub(r"\W+", "-", p.rsplit("/", 1)[-1].rsplit(".", 1)[0]).strip("-")
    keep = []
    try:
        for line in open(out, encoding="utf-8"):
            line = line.rstrip("\n")
            if line and not line.startswith("#") and not any(f"kit={tag(k)} " in line + " " for k in kits):
                keep.append(line)
    except FileNotFoundError:
        pass
    tests, skipped = [], []
    for k in kits:
        for n, bn, typed, who in kit_rows(k):
            b, t = words(bn), [w.lower() for w in words(typed)]
            if len(b) != len(t):
                skipped.append(f"{tag(k)} #{n}: {len(t)} typed words for {len(b)} Bangla words")
                continue
            for i in range(len(b)):
                tests.append(f"{' '.join(t[:i + 1])}\t{b[i]}\tkit={tag(k)} #{n} {who}".rstrip())
    with open(out, "w", encoding="utf-8") as f:
        f.write("# HELD-OUT: typed by people outside the project with the typing kit (scripts/typing_kit.py).\n"
                "# Never tune on these; only the total is looked at. Last word scored, the words before it are context.\n"
                "# typed\taccepted\tsource\n")
        for line in dict.fromkeys(keep + tests):
            f.write(line + "\n")
    print(f"{len(tests)} tests from {len(kits)} kit(s); {len(keep)} kept from before")
    for s in skipped:
        print("skipped", s)


if __name__ == "__main__":
    main()
