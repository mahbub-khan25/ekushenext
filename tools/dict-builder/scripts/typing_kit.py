#!/usr/bin/env python3
"""Make the typing kit: everyday Bangla sentences for a friend to type the way they normally would,
in English letters. import_typing.py turns the filled-in kit into a held-out test set
(eval/community.tsv): every typed word is a test, with the words before it as context.

Usage: typing_kit.py OUT.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

# Standard spelling (প্রমিত বানান), as Ekushe should write them; written for this kit, not taken from
# Ekushe's other test sets or chat phrases.
SENTENCES = [
    # greetings and small talk
    "আসসালামু আলাইকুম, কেমন আছেন?", "আমি ভালো আছি, আপনি কেমন আছেন?", "অনেক দিন পর দেখা হলো।",
    "তোমার শরীর এখন কেমন?", "বাসার সবাই ভালো আছে তো?", "আজকে আবহাওয়াটা খুব সুন্দর।",
    "কাল রাতে অনেক বৃষ্টি হয়েছে।", "আজ খুব গরম পড়েছে।", "শুভ সকাল, দিনটা ভালো কাটুক।",
    "শুভ রাত্রি, ভালো করে ঘুমাও।",
    # family
    "আম্মু রান্না করছে, একটু পরে খেতে বসব।", "আব্বু অফিস থেকে এখনো ফেরেনি।", "ছোট ভাইয়ের কাল পরীক্ষা।",
    "আপুর বিয়ে আগামী মাসে।", "নানুর বাড়ি যাব ঈদের ছুটিতে।", "বাচ্চাটা সারাদিন খুব দুষ্টুমি করেছে।",
    "মামা আজকে ঢাকায় আসবেন।", "আমার বোন এবার বিশ্ববিদ্যালয়ে ভর্তি হয়েছে।",
    # food
    "দুপুরে কী খেয়েছ?", "আজ রাতে খিচুড়ি আর ডিম ভাজি হবে।", "চা খাবে নাকি কফি?",
    "বিরিয়ানিটা অসাধারণ হয়েছে।", "রেস্টুরেন্টে খুব ভিড় ছিল।", "ফুচকা খেতে ইচ্ছা করছে।", "মিষ্টি একটু কম দিও।",
    # work and study
    "আজকে অফিসে অনেক কাজের চাপ।", "মিটিংটা বিকেল চারটায় শুরু হবে।", "রিপোর্টটা কালকের মধ্যে জমা দিতে হবে।",
    "বস আজকে খুব খুশি ছিলেন।", "ক্লাস শেষ হতে দেরি হবে।", "পরীক্ষার প্রস্তুতি কেমন চলছে?",
    "আমি এখনো অ্যাসাইনমেন্টটা শেষ করতে পারিনি।", "স্যার আজকে ক্লাসে আসেননি।", "নতুন চাকরিটা কেমন লাগছে?",
    "ছুটির আবেদন করেছি, এখনো উত্তর পাইনি।",
    # getting around
    "জ্যামে আটকে আছি, আসতে দেরি হবে।", "বাসে উঠে তোমাকে ফোন দেব।", "আমরা শুক্রবার কক্সবাজার যাচ্ছি।",
    "ট্রেনের টিকিট পাওয়া যায়নি।", "রিকশা ভাড়া অনেক বেড়ে গেছে।", "উবারে করে চলে আসো।",
    "চট্টগ্রাম থেকে কখন রওনা দেবে?", "সিলেটে এখন খুব ঠান্ডা।", "মিরপুর দশে নেমে আমাকে কল দিও।",
    # shopping and money
    "কত টাকা লাগবে বলো।", "বিকাশে টাকা পাঠিয়ে দিয়েছি।", "দামটা একটু বেশি মনে হচ্ছে।",
    "নতুন মোবাইলটা কিনেছি গতকাল।", "অনলাইনে অর্ডার করলে তাড়াতাড়ি পাবে।", "বাজার থেকে কিছু আনতে হবে?",
    # phone and online
    "তোমার মেসেজ দেখিনি, দুঃখিত।", "ছবিগুলো হোয়াটসঅ্যাপে পাঠাও।", "ফেসবুকে পোস্টটা দেখেছ?",
    "নেট খুব স্লো, ভিডিও কল করা যাচ্ছে না।", "ফোনে চার্জ নেই, পরে কথা বলি।", "ইউটিউবে নতুন ভিডিওটা দেখো।",
    "লিংকটা আমাকে দাও তো।",
    # health
    "মাথা ব্যথা করছে, একটু বিশ্রাম নেব।", "ডাক্তারের কাছে গিয়েছিলাম আজকে।", "জ্বর এখন একটু কমেছে।",
    "ওষুধটা ঠিক সময়ে খেয়ো।", "নিজের যত্ন নিও।",
    # feelings
    "আজ মনটা খুব ভালো লাগছে।", "তোমার কথা খুব মনে পড়ছে।", "খুব টেনশনে আছি পরীক্ষা নিয়ে।",
    "চিন্তা কোরো না, সব ঠিক হয়ে যাবে।", "তুমি না থাকলে কী যে করতাম।", "অনেক ধন্যবাদ তোমার সাহায্যের জন্য।",
    "রাগ কোরো না প্লিজ।", "শুনে খুব খারাপ লাগল।",
    # plans
    "কাল সন্ধ্যায় দেখা করি?", "শুক্রবার কি তুমি ফ্রি আছ?", "আমরা সবাই মিলে ঘুরতে যাব।",
    "সিনেমাটা দেখতে যাবে?", "ঈদের ছুটিতে বাড়ি যাচ্ছি।", "জন্মদিনের পার্টি কোথায় হবে?",
    "পরে জানাচ্ছি, একটু অপেক্ষা করো।",
    # questions
    "তুমি এখন কোথায়?", "কখন আসবে বলো তো?", "এটা কার বই?", "কী হয়েছে, এত চুপচাপ কেন?",
    "তুমি কি আমার কথা শুনতে পাচ্ছ?", "এর দাম কত?", "তোমার নাম্বারটা দেবে?",
    # written Bangla
    "সরকার নতুন বাজেট ঘোষণা করেছে।", "দেশের অর্থনীতি এখন কঠিন সময় পার করছে।", "বাংলাদেশ ক্রিকেট দল আজ জিতেছে।",
    "শিক্ষার্থীদের জন্য নতুন বৃত্তি চালু হয়েছে।", "ঢাকায় যানজট দিন দিন বাড়ছে।",
    "একুশে ফেব্রুয়ারি আন্তর্জাতিক মাতৃভাষা দিবস।",
    # chat
    "কী করছ এখন?", "ভাত খেয়েছ?", "আমি বাসায় আছি।", "একটু পরে ফোন দিচ্ছি।", "সমস্যা নেই, আমি সামলে নেব।",
    "ঠিক আছে, দেখা হবে।", "হ্যাঁ, আমি রাজি।", "না ভাই, আজকে পারব না।", "আরে, তুমি তো দারুণ লিখেছ!",
    "আচ্ছা, তাহলে কাল কথা হবে।",
]

LATIN, BANGLA = "Arial", "Nirmala UI"
HEAD_FILL = PatternFill("solid", fgColor="1F4E46")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="C9C9C9")


def main():
    out = sys.argv[1]
    wb = Workbook()
    how = wb.active
    how.title = "How to fill"
    kit = wb.create_sheet("Sentences")

    heads = ["#", "Bangla sentence", "How you type it (English letters)", "Note (optional)"]
    for c, (h, w) in enumerate(zip(heads, [6, 46, 52, 28]), 1):
        cell = kit.cell(row=1, column=c, value=h)
        cell.font = Font(name=LATIN, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        kit.column_dimensions[cell.column_letter].width = w
    kit.freeze_panes = "C2"
    rows = [["Example", "আমি কাল ঢাকা যাচ্ছি।", "ami kal dhaka jacchi", "how I'd type it in a chat"]]
    rows += [[i, s, None, None] for i, s in enumerate(SENTENCES, 1)]
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = kit.cell(row=r, column=c, value=v)
            cell.font = Font(name=BANGLA if c == 2 else LATIN, size=13 if c == 2 else 11,
                             italic=(r == 2), color="555555" if r == 2 else "000000")
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if r == 2:
                cell.fill = EXAMPLE_FILL
            elif c in (3, 4):
                cell.fill = INPUT_FILL

    how.column_dimensions["A"].width = 30
    how.column_dimensions["B"].width = 72
    how["A1"] = "Ekushe Next — typing kit"
    how["A1"].font = Font(name=LATIN, bold=True, size=16, color="1F4E46")
    lines = [
        ("What this is", f"{len(SENTENCES)} everyday Bangla sentences. Type each one in English letters, the way you normally type Bangla in a chat (Banglish). It shows Ekushe how real people type."),
        ("What to fill in", "The yellow cells only: on the Sentences sheet, 'How you type it' for each sentence; below, a few words about you."),
        ("How to type", "Type naturally, as fast as you would in a message: no need to be careful or 'correct'. Keep one typed word for each Bangla word, in the same order. Punctuation doesn't matter."),
        ("Skipping", "Leave a sentence empty if you'd never write it. A few is plenty; all of them is wonderful."),
        ("Row 2", "The grey row is an example; it is ignored."),
        ("When done", "Save the file (or download it from Google Sheets as .xlsx) and send it back. Nothing you type here is published with your name."),
    ]
    for i, (k, v) in enumerate(lines, 3):
        how.cell(row=i, column=1, value=k).font = Font(name=LATIN, bold=True)
        c = how.cell(row=i, column=2, value=v)
        c.font = Font(name=LATIN)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    p = len(lines) + 4
    how.cell(row=p, column=1, value="About you").font = Font(name=LATIN, bold=True, size=12, color="1F4E46")
    about = [("Your first name (optional)", ""), ("District you grew up in", "e.g. Dhaka, Chattogram, Sylhet, Rajshahi"),
             ("Keyboard you use now", "")]
    for i, (k, hint) in enumerate(about, p + 1):
        how.cell(row=i, column=1, value=k).font = Font(name=LATIN)
        c = how.cell(row=i, column=2, value=None)
        c.fill = INPUT_FILL
        c.font = Font(name=LATIN)
        if hint:
            how.cell(row=i, column=3, value=hint).font = Font(name=LATIN, italic=True, color="777777")
    how.column_dimensions["C"].width = 40
    dv = DataValidation(type="list", formula1='"Avro,Bijoy,Gboard,Ridmik,Samsung keyboard,Ekushe Next,Other"', allow_blank=True)
    how.add_data_validation(dv)
    dv.add(f"B{p + 3}")
    wb.save(out)
    print(f"{out}: {len(SENTENCES)} sentences")


if __name__ == "__main__":
    main()
