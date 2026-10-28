#!/usr/bin/env python3
"""Bring answers from the spelling-review workbook (review_sheet.py) into curated/spelling-seed.tsv.

Reads the Review sheet of an .xlsx, or a .csv saved from it (Google Sheets: File → Download). Each
answered row becomes a reviewed decision; a pair already in the seed file is replaced, so answering
again corrects an earlier answer. 'Not sure' and empty rows are skipped.

Usage: import_review.py FILLED.xlsx|.csv curated/spelling-seed.tsv [REVIEWER]
"""
import csv
import sys
import unicodedata

B_RIGHT, A_RIGHT, BOTH = ("B is right (A is a misspelling)", "A is right (B is a misspelling)",
                          "Both right (different words)")


def rows_of(path):
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig") as f:
            r = list(csv.reader(f))
    else:
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True)["Review"]
        r = [[("" if v is None else str(v)) for v in row] for row in ws.iter_rows(values_only=True)]
    head = [h.strip() for h in r[0]]
    col = {h: i for i, h in enumerate(head)}
    for row in r[1:]:
        get = lambda h: row[col[h]].strip() if h in col and col[h] < len(row) else ""
        if get("#") in ("", "Example"):
            continue
        yield get("Word A"), get("Word B (Ekushe's standard)"), get("Your answer"), get("Note (optional)")


def main():
    src, seed = sys.argv[1], sys.argv[2]
    who = sys.argv[3] if len(sys.argv) > 3 else "reviewer"
    nfc = lambda s: unicodedata.normalize("NFC", s)
    decisions = {}
    for a, b, ans, note in rows_of(src):
        a, b = nfc(a), nfc(b)
        if not a or not b:
            continue
        note = " ".join(f"{note} (reviewed by {who})".split()).replace("\t", " ")
        if ans == B_RIGHT:
            decisions[frozenset((a, b))] = f"{b}\t{a}\tprefer\treviewed\t{note}"
        elif ans == A_RIGHT:
            decisions[frozenset((a, b))] = f"{a}\t{b}\tprefer\treviewed\t{note}"
        elif ans == BOTH:
            decisions[frozenset((a, b))] = f"{a}\t{b}\tdistinct\treviewed\t{note}"
    kept, replaced = [], 0
    for line in open(seed, encoding="utf-8"):
        line = line.rstrip("\n")
        f = line.split("\t")
        if not line.startswith("#") and len(f) == 5:
            words = {nfc(f[0]), *(nfc(v) for v in f[1].split(",") if v)}
            if len(words) == 2 and frozenset(words) in decisions:
                replaced += 1
                continue
        kept.append(line)
    if decisions:
        kept.append(f"# ── reviewed with the spelling-review workbook ({who})")
        kept.extend(decisions.values())
    with open(seed, "w", encoding="utf-8") as f:
        f.write("\n".join(kept) + "\n")
    print(f"{len(decisions)} decisions written ({replaced} replaced earlier ones)")


if __name__ == "__main__":
    main()
